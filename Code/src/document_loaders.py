"""Document loading for Markdown, text, PDF, and DOCX source material."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from .utils import utc_timestamp


SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}


@dataclass(frozen=True)
class LoadedDocument:
    """A document or page loaded from local study material."""

    text: str
    metadata: dict[str, str | int]


def _base_metadata(path: Path, subject_key: str) -> dict[str, str]:
    """Build metadata shared by all document types."""
    return {
        "subject": subject_key,
        "source_path": str(path),
        "source_name": path.name,
        "source_type": path.suffix.lower().lstrip("."),
        "loaded_at": utc_timestamp(),
    }


def load_text_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a Markdown or plain text file as one document."""
    text = path.read_text(encoding="utf-8", errors="replace")
    return [LoadedDocument(text=text, metadata=_base_metadata(path, subject_key))]


def load_pdf_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a PDF file page by page using pypdf when it is installed."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on optional package
        raise RuntimeError("pypdf is required to load PDF files.") from exc

    documents: list[LoadedDocument] = []
    reader = PdfReader(str(path))
    for page_index, page in enumerate(reader.pages, start=1):
        metadata = _base_metadata(path, subject_key)
        metadata["page_number"] = page_index
        documents.append(LoadedDocument(text=page.extract_text() or "", metadata=metadata))
    return documents


def load_docx_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a DOCX file using python-docx when it is installed."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on optional package
        raise RuntimeError("python-docx is required to load DOCX files.") from exc

    docx_file = Document(str(path))
    paragraphs = [paragraph.text for paragraph in docx_file.paragraphs if paragraph.text.strip()]
    return [LoadedDocument(text="\n".join(paragraphs), metadata=_base_metadata(path, subject_key))]


def load_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load one supported file and attach source metadata."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return load_text_document(path, subject_key)
    if suffix == ".pdf":
        return load_pdf_document(path, subject_key)
    if suffix == ".docx":
        return load_docx_document(path, subject_key)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def iter_material_files(folders: Iterable[Path]) -> list[Path]:
    """Return supported material files from the given folders."""
    files: list[Path] = []
    for folder in folders:
        if folder.exists():
            # Sorting makes indexing predictable for testing and debugging.
            files.extend(sorted(path for path in folder.rglob("*") if path.suffix.lower() in SUPPORTED_EXTENSIONS))
    return files


def load_subject_documents(subject_key: str, folders: Iterable[Path]) -> list[LoadedDocument]:
    """Load all supported source files for a subject."""
    documents: list[LoadedDocument] = []
    for path in iter_material_files(folders):
        documents.extend(load_document(path, subject_key))
    return documents

