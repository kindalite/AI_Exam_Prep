"""Document loading for Markdown, text, PDF, DOCX, and local image material."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from .utils import utc_timestamp


SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".svg"}


@dataclass(frozen=True)
class LoadedDocument:
    """A document or page loaded from local study material."""

    text: str
    metadata: dict[str, str | int]


def _base_metadata(path: Path, subject_key: str) -> dict[str, str]:
    """Build metadata shared by all document types."""
    return {
        "subject": subject_key,
        "source_path": str(path),
        "source_name": path.name,
        "source_type": path.suffix.lower().lstrip("."),
        "source_layer": "local_material",
        "loaded_at": utc_timestamp(),
    }


def load_text_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a Markdown or plain text file as one document."""
    text = path.read_text(encoding="utf-8", errors="replace")
    return [LoadedDocument(text=text, metadata=_base_metadata(path, subject_key))]


def load_pdf_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a PDF file page by page using pypdf when it is installed."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on optional package
        raise RuntimeError("pypdf is required to load PDF files.") from exc

    documents: list[LoadedDocument] = []
    reader = PdfReader(str(path))
    for page_index, page in enumerate(reader.pages, start=1):
        metadata = _base_metadata(path, subject_key)
        metadata["page_number"] = page_index
        metadata["modality"] = "pdf_text"
        documents.append(LoadedDocument(text=page.extract_text() or "", metadata=metadata))
    return documents


def load_docx_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Load a DOCX file using python-docx when it is installed."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on optional package
        raise RuntimeError("python-docx is required to load DOCX files.") from exc

    docx_file = Document(str(path))
    paragraphs = [paragraph.text for paragraph in docx_file.paragraphs if paragraph.text.strip()]
    return [LoadedDocument(text="\n".join(paragraphs), metadata=_base_metadata(path, subject_key))]


def load_svg_document(path: Path, subject_key: str) -> list[LoadedDocument]:
    """Safely load readable SVG text and metadata without executing content."""
    import xml.etree.ElementTree as ET

    raw = path.read_text(encoding="utf-8", errors="replace")
    metadata = _base_metadata(path, subject_key)
    metadata["modality"] = "svg_text"
    lowered = raw.lower()
    if "<script" in lowered or "javascript:" in lowered:
        metadata["svg_warning"] = "Unsafe SVG script content was ignored during indexing."
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        metadata["svg_warning"] = f"SVG metadata parse failed: {exc}"
        return [LoadedDocument(text="", metadata=metadata)]
    texts: list[str] = []
    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1].lower()
        if tag in {"title", "desc", "text", "tspan"} and element.text and element.text.strip():
            texts.append(element.text.strip())
    return [LoadedDocument(text="\n".join(texts), metadata=metadata)]


def load_image_document(path: Path, subject_key: str, config=None) -> list[LoadedDocument]:
    """Load an image as OCR and local vision text when possible."""
    from .config import load_config
    from .image_understanding import describe_image_safe
    from .ocr import ocr_image_safe

    app_config = config or load_config()
    documents: list[LoadedDocument] = []
    if getattr(app_config, "enable_ocr", True):
        text, warning = ocr_image_safe(path, getattr(app_config, "ocr_languages", "deu+eng+fra"))
        metadata = _base_metadata(path, subject_key)
        metadata["modality"] = "ocr_text"
        if warning:
            metadata["ocr_warning"] = warning
        documents.append(LoadedDocument(text=text, metadata=metadata))
    if getattr(app_config, "enable_image_understanding", True):
        description, warning = describe_image_safe(path, subject_key, app_config)
        metadata = _base_metadata(path, subject_key)
        metadata["modality"] = "image_description"
        if warning:
            metadata["vision_warning"] = warning
        documents.append(LoadedDocument(text=description, metadata=metadata))
    return documents


def load_document(path: Path, subject_key: str, config=None) -> list[LoadedDocument]:
    """Load one supported file and attach source metadata."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return load_text_document(path, subject_key)
    if suffix == ".pdf":
        if config is not None and getattr(config, "enable_pdf_page_rendering", True):
            try:
                from .multimodal_pdf import load_pdf_multimodal

                return load_pdf_multimodal(path, subject_key, config)
            except Exception:
                return load_pdf_document(path, subject_key)
        return load_pdf_document(path, subject_key)
    if suffix == ".docx":
        return load_docx_document(path, subject_key)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return load_image_document(path, subject_key, config=config)
    if suffix == ".svg":
        return load_svg_document(path, subject_key)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def iter_material_files(folders: Iterable[Path]) -> list[Path]:
    """Return supported material files from the given folders."""
    files: list[Path] = []
    for folder in folders:
        if folder.exists() and folder.is_file() and folder.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(folder)
        elif folder.exists():
            files.extend(sorted(path for path in folder.rglob("*") if path.suffix.lower() in SUPPORTED_EXTENSIONS))
    return files


def load_subject_documents(subject_key: str, folders: Iterable[Path], config=None) -> list[LoadedDocument]:
    """Load all supported source files for a subject."""
    documents: list[LoadedDocument] = []
    for path in iter_material_files(folders):
        documents.extend(load_document(path, subject_key, config=config))
    return documents
